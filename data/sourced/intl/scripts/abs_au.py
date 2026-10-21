"""Australian Bureau of Statistics: Counts of Australian Businesses, including Entries and Exits, June 2021 to June 2025
(8165.0, released 16 December 2025: data cubes 2, 5 and 10), Australian Industry 2024-25 (8155.0, ABS Data API dataflow
AUSTRALIAN_INDUSTRY 1.1.0), Estimated resident population by LGA (2025), 30 June 2025 (ABS Data API ERP_LGA2025).
Downloads in /srv/atlas/runs/data-intl/dl/abs. Run with the venv python (openpyxl).
"""
import csv
import openpyxl
from lib import write, DL, RETRIEVED

D = f"{DL}/abs"
CABEE = "https://www.abs.gov.au/statistics/economy/business-indicators/counts-australian-businesses-including-entries-and-exits/jul2021-jun2025"
SRC = {
    "dc02": {"publisher": "Australian Bureau of Statistics", "title": "Counts of Australian Businesses, including Entries and Exits, June 2021 to June 2025, data cube 2: businesses by main state by industry class by employment size ranges", "url": CABEE + "/8165DC02.xlsx"},
    "dc05": {"publisher": "Australian Bureau of Statistics", "title": "Counts of Australian Businesses, including Entries and Exits, June 2021 to June 2025, data cube 5: survival of entries by main state by subdivision by employment size ranges", "url": CABEE + "/8165DC05.xlsx"},
    "dc10": {"publisher": "Australian Bureau of Statistics", "title": "Counts of Australian Businesses, including Entries and Exits, June 2021 to June 2025, data cube 10: businesses by industry division by Local Government Area by employment size ranges", "url": CABEE + "/8165DC10.xlsx"},
    "ai": {"publisher": "Australian Bureau of Statistics", "title": "Australian Industry, 2024-25 (ABS Data API, AUSTRALIAN_INDUSTRY)", "url": "https://www.abs.gov.au/statistics/industry/industry-overview/australian-industry/latest-release"},
    "erp": {"publisher": "Australian Bureau of Statistics", "title": "Estimated resident population by LGA (2025), 30 June 2025 (ABS Data API, ERP_LGA2025)", "url": "https://www.abs.gov.au/statistics/people/population/regional-population/latest-release"},
}
PERT = "The ABS perturbs small cells in this table (totals of industry division, state and Australia excepted): treat small counts as approximate."
AUS = {"level": "country", "code": "AU", "name": "Australia"}
STATES = {"New South Wales": ("AU-NSW", "1"), "Victoria": ("AU-VIC", "2"), "Queensland": ("AU-QLD", "3"), "South Australia": ("AU-SA", "4"),
          "Western Australia": ("AU-WA", "5"), "Australian Capital Territory": ("AU-ACT", "8")}  # the states of the six catalog cities


def sgeo(name):
    return {"level": "region", "code": STATES[name][0], "name": name}


rows = []

# 1. Business counts by class, June 2025 (operating at end of financial year 2024-25), Australia and states
CLASSES = {4511: (["restaurants", "cafes"], "ANZSIC 4511 Cafes and Restaurants is one class: cafes and restaurants are counted together."),
           4520: ([None], "Pubs, taverns and bars."),
           9511: (["barbershops"], "ANZSIC 9511 Hairdressing and Beauty Services: barbers, hairdressers and beauty salons together.")}
BANDS = {27: "no employees (non-employing)", 28: "1 to 4 employees", 29: "5 to 19 employees", 30: "20 to 199 employees", 31: "200 or more employees"}
ws = openpyxl.load_workbook(f"{D}/8165DC02.xlsx", read_only=True)["Table 1"]
for r in ws.iter_rows(values_only=True):
    geo = cls = None
    if r[0] in STATES and r[1] in CLASSES:
        geo, cls, lab = sgeo(r[0]), r[1], r[2]
    elif r[0] is None and r[1] in CLASSES and isinstance(r[2], str) and r[2].startswith("Total "):
        geo, cls, lab = AUS, r[1], r[2][6:]
    elif r[0] == "Total Australia":
        geo, cls, lab = AUS, "all", None
    if not geo:
        continue
    off = 0  # total rows keep empty code and label cells: same columns as the class rows
    total = r[32]
    trades, note = (["all"], None) if cls == "all" else CLASSES[cls]
    oc = "ANZSIC 2006, all industries (market sector)" if cls == "all" else f"ANZSIC 2006 {cls} {lab}"
    for t in trades:
        rows.append({k: v for k, v in dict(item="business_count", trade=t, officialClass=oc, geo=geo, value=int(total), unit="businesses",
                     year="2025-06", measure="actively trading businesses (ABN-registered, GST-registered market sector) operating at the end of the 2024-25 financial year, by main state",
                     note=" ".join(x for x in (note, PERT if cls != "all" else None) if x), source=SRC["dc02"]).items() if v not in (None, "")})
        if geo is AUS:
            for col, band in BANDS.items():
                n = r[col + off]
                rows.append({k: v for k, v in dict(item="size_class_share", trade=t, officialClass=oc, geo=geo, value=round(100 * n / total, 1), unit="percent of businesses",
                             year="2025-06", measure=f"share of businesses with {band} (annualised employment)", note=f"{n:,} of {int(total):,} businesses. " + PERT if cls != "all" else f"{n:,} of {int(total):,} businesses.",
                             source=SRC["dc02"]).items() if v not in (None, "")})

# 2. Survival of the businesses that entered in 2021-22, to June 2023, 2024 and 2025, by subdivision, states and Australia
SUBS = {45: (["restaurants", "cafes"], "ANZSIC subdivision 45, all food and beverage services (cafes and restaurants, takeaway, catering, pubs and bars, clubs)."),
        95: (["barbershops"], "ANZSIC subdivision 95, all personal and other services (hairdressing and beauty, diet and weight, laundry, funeral, religious and other services).")}
ws = openpyxl.load_workbook(f"{D}/8165DC05.xlsx", read_only=True)["Table 1"]
for r in ws.iter_rows(values_only=True):
    if r[0] in STATES and r[1] in SUBS:
        geo, sub, trades, note, v = sgeo(r[0]), r[1], SUBS[r[1]][0], SUBS[r[1]][1], r
        oc = f"ANZSIC 2006 subdivision {r[1]} {r[2]}"
    elif r[0] == "Total Australia":
        geo, trades, note, v, oc = AUS, ["all"], None, r, "ANZSIC 2006, all industries (market sector)"
    elif isinstance(r[0], str) and r[0].startswith("Total ") and r[0][6:] in STATES:
        geo, trades, note, v, oc = sgeo(r[0][6:]), ["all"], None, r, "ANZSIC 2006, all industries (market sector)"
    else:
        continue
    entries = v[8]
    for n, (cnt, rate, yr) in ((1, (19, 20, "2023-06")), (2, (31, 32, "2024-06")), (3, (43, 44, "2025-06"))):
        for t in trades:
            rows.append({k: x for k, x in dict(item=f"survival_{n}y", trade=t, officialClass=oc, geo=geo, value=v[rate], unit="percent of businesses that entered",
                         year=yr, measure=f"share of the businesses that entered in 2021-22 still trading in June {yr[:4]} (about {n} year{'s' if n > 1 else ''} after entry; all businesses, employing or not)",
                         note=" ".join(z for z in (note, f"{v[cnt]:,} of {entries:,} entries survived (published rate).", PERT if trades != ['all'] else None) if z),
                         source=SRC["dc05"]).items() if x not in (None, "")})

# 3. Business counts by LGA for the six capital-city councils, June 2025, by division (no class is published by LGA)
LGAS = {17200: ("Sydney", "City of Sydney (local government area)"), 24600: ("Melbourne", "City of Melbourne (local government area)"),
        31000: ("Brisbane", "City of Brisbane (local government area)"), 40070: ("Adelaide", "City of Adelaide (local government area)"),
        57080: ("Perth", "City of Perth (local government area)"), 89399: ("Unincorporated ACT", "Australian Capital Territory (Canberra; unincorporated, one LGA)")}
DIV = {"H": (None, "Whole division H, accommodation and food services: hotels as well as cafes, restaurants, takeaway and pubs. No class is published by LGA."),
       "S": (None, "Whole division S, other services: repair and maintenance as well as personal services. No class is published by LGA.")}
LGANOTE = ("Businesses that operate in several states are placed at state level only and are not in any LGA count; "
           "counts are by the business's main address.")
ws = openpyxl.load_workbook(f"{D}/8165DC10.xlsx", read_only=True)["Table 1"]
tot = {}
for r in ws.iter_rows(values_only=True):
    if r[1] in LGAS and isinstance(r[3], str):
        tot[r[1]] = tot.get(r[1], 0) + int(r[10])
        if r[3] in DIV:
            rows.append(dict(item="business_count", officialClass=f"ANZSIC 2006 division {r[3]} {r[4]}", geo={"level": "city", "code": f"AU-LGA-{r[1]}", "name": LGAS[r[1]][1]},
                             value=int(r[10]), unit="businesses", year="2025-06", measure="actively trading businesses at June 2025, by the main address's local government area",
                             note=" ".join((DIV[r[3]][1], LGANOTE, PERT.replace("(totals of industry division, state and Australia excepted)", "(division totals excepted)"))), source=SRC["dc10"]))
for code, n in tot.items():
    rows.append(dict(item="business_count", trade="all", officialClass="ANZSIC 2006, all industries (market sector), sum of the published divisions",
                     geo={"level": "city", "code": f"AU-LGA-{code}", "name": LGAS[code][1]}, value=n, unit="businesses", year="2025-06",
                     measure="actively trading businesses at June 2025, by the main address's local government area: the sum of the published industry-division totals (the table prints no all-industry row)",
                     note=LGANOTE, source=SRC["dc10"]))

# 4. Populations, 30 June 2025: Australia, the six states, the six LGAs
for r in csv.DictReader(open(f"{D}/erp_lga2025.csv", encoding="utf-8-sig")):
    rt, reg = r["REGION_TYPE: Region Type"].split(":")[0], r["REGION: Region"]
    code, name = reg.split(": ", 1)
    if r["TIME_PERIOD: Time Period"] != "2025" or not r["MEASURE: Measure"].startswith("ERP"):
        continue
    if rt == "AUS":
        g = AUS
    elif rt == "STE" and name in STATES:
        g = sgeo(name)
    elif rt.startswith("LGA") and code.isdigit() and int(code) in LGAS:
        g = {"level": "city", "code": f"AU-LGA-{code}", "name": LGAS[int(code)][1]}
    else:
        continue
    if any(x["item"] == "population" and x["geo"]["code"] == g["code"] for x in rows):
        continue
    rows.append(dict(item="population", geo=g, value=int(float(r["OBS_VALUE"])), unit="persons", year="2025-06",
                     measure="estimated resident population, 30 June 2025", note="LGA boundaries of 2025; the business counts' LGA codes are the same." if g["level"] == "city" else "Same boundary as the business counts.",
                     source=SRC["erp"]))

# 5. Australian Industry 2024-25: subdivisions 45 and 95 and the total of selected industries, Australia
ai = {}
for r in csv.DictReader(open(f"{D}/ai.csv", encoding="utf-8-sig")):
    if r["REGION: Region"].startswith("AUS") and r["TIME_PERIOD: Time Period"] == "2025":
        ai[(r["INDUSTRY: Industry"].split(":")[0], r["MEASURE: Measure"].split(":")[0])] = (float(r["OBS_VALUE"]), r["OBS_STATUS: Observation Status"])
AI = {"45": (["restaurants", "cafes"], "ANZSIC 2006 subdivision 45 Food and Beverage Services", SUBS[45][1]),
      "95": (["barbershops"], "ANZSIC 2006 subdivision 95 Personal and Other Services", SUBS[95][1]),
      "TOT_NO_K": (["all"], "ANZSIC 2006, total selected industries (excluding K, financial and insurance services)", None)}
for ind, (trades, oc, scope) in AI.items():
    s = ai.get((ind, "INCSALGDSSERV"))
    if not s:
        continue
    for m, item, measure in (("LABCOSTTOT", "staff_cost_share", "total labour costs (wages and salaries, superannuation, workers' compensation, payroll tax, fringe benefits tax) divided by sales and service income"),
                             ("PURCHASES", "purchases_share", "purchases of goods and materials (not services, not rent) divided by sales and service income"),
                             ("EBITDA", "ebitda_share", "earnings before interest, tax, depreciation and amortisation divided by sales and service income (EBITDA, the ABS's measure; close to but not the same as a gross operating surplus)"),
                             ("OPBT", "operating_profit_margin", "operating profit before tax divided by sales and service income (after depreciation and all pay)")):
        v = ai.get((ind, m))
        if not v:
            continue
        st = v[1]
        if st.startswith(("z", "w")):  # too unreliable for general use: not written
            continue
        q = f"ABS quality flag: {st.split(': ', 1)[1]}." if st else None
        for t in trades:
            rows.append({k: x for k, x in dict(item=item, trade=t, officialClass=oc, geo=AUS, value=round(100 * v[0] / s[0], 1), unit="percent of sales and service income",
                         year="2024-25", measure=measure, note=" ".join(z for z in (scope, f"{v[0]:,.0f} million AUD over sales and service income of {s[0]:,.0f} million AUD.", q) if z),
                         source=SRC["ai"]).items() if x not in (None, "")})

write("abs-australia.json", {
    "id": "abs-australia",
    "source": {"publisher": "Australian Bureau of Statistics", "title": "Counts of Australian Businesses (June 2025), Australian Industry (2024-25), Regional population (30 June 2025)",
               "url": CABEE, "table": "Each row names its own data cube or dataflow: 8165DC02, 8165DC05, 8165DC10, AUSTRALIAN_INDUSTRY, ERP_LGA2025",
               "licence": "Creative Commons Attribution 4.0 International (ABS copyright: www.abs.gov.au/website-privacy-copyright-and-disclaimer)"},
    "retrieved": RETRIEVED, "rows": rows})

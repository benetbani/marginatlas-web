"""Japan: 2021 Economic Census for Business Activity (令和3年経済センサス‐活動調査, Statistics Bureau of Japan and METI),
tables 7-1 (national, by JSIC minor group and size: establishments, persons engaged, sales, sales per establishment) and
9-1A (all establishments by minor group, municipalities); 2020 Population Census table 1-1 (population by municipality).
Files from e-Stat (https://www.e-stat.go.jp/stat-search/file-download?statInfId=<id>&fileKind=0) in /srv/atlas/runs/data-intl/dl/jp.
Run with the venv python (openpyxl).
"""
import openpyxl
from lib import write, DL, RETRIEVED

D = f"{DL}/jp"
DLURL = "https://www.e-stat.go.jp/stat-search/file-download?statInfId={}&fileKind=0"
SRC = {
    "t71": {"publisher": "Statistics Bureau of Japan (Ministry of Internal Affairs and Communications) and Ministry of Economy, Trade and Industry",
            "title": "2021 Economic Census for Business Activity, establishments, cross-industry tabulation, table 7-1 (JSIC minor group by size of persons engaged: establishments, persons engaged, sales, sales per establishment), Japan",
            "url": DLURL.format("000040067950")},
    "t91": {"publisher": "Statistics Bureau of Japan (Ministry of Internal Affairs and Communications) and Ministry of Economy, Trade and Industry",
            "title": "2021 Economic Census for Business Activity, establishments, cross-industry tabulation, table 9-1A (all establishments by JSIC minor group, Japan, prefectures, municipalities)",
            "url": DLURL.format("000040067884")},
    "pop": {"publisher": "Statistics Bureau of Japan (Ministry of Internal Affairs and Communications)",
            "title": "2020 Population Census, basic complete tabulation on population and households, table 1-1 (population by municipality)",
            "url": DLURL.format("000032142402")},
}
CLASSES = {  # JSIC Rev. 13 code: (trade, note)
    "AR": ("all", None),
    "76": ("restaurants", "JSIC group 76, eating and drinking places: restaurants of every kind together with coffee shops (767) and bars (765, 766). Japan has no single class for restaurants: the general and specialty restaurant classes (761 to 764) are written apart."),
    "761": (None, "JSIC 761: canteens and restaurants other than specialty restaurants."),
    "762": (None, "JSIC 762: specialty restaurants (Japanese, Chinese, Western, grilled meat and other cuisine)."),
    "763": (None, None), "764": (None, None), "765": (None, None), "766": (None, None), "769": (None, None),
    "767": ("cafes", "JSIC 767: coffee shops (喫茶店), the class for cafes."),
    "782": ("barbershops", "JSIC 782: barber shops (理容業) alone; beauty salons are 783."),
    "783": (None, "JSIC 783: beauty salons (美容業)."),
}
JP = {"level": "country", "code": "JP", "name": "Japan"}
CITIES = {"01100": "Sapporo", "04100": "Sendai", "13100": "Tokyo (the 23 special wards)", "14100": "Yokohama", "23100": "Nagoya",
          "26100": "Kyoto", "27100": "Osaka", "28100": "Kobe", "34100": "Hiroshima", "40130": "Fukuoka"}


def cg(code):
    return {"level": "city", "code": f"JP-{code}", "name": CITIES[code]}


rows = []
BANDS = {"1_1～4人": "1 to 4", "2_5～9人": "5 to 9", "3_10～19人": "10 to 19", "4_20～29人": "20 to 29", "5_30～49人": "30 to 49", "6_50人以上": "50 or more"}

# 1. National: table 7-1
ws = openpyxl.load_workbook(f"{D}/000040067950.xlsx", read_only=True).active
for r in ws.iter_rows(values_only=True):
    if r[0] != "00_全国" or not isinstance(r[3], str):
        continue
    code, label = r[3].split("_", 1)
    if code not in CLASSES:
        continue
    trade, note = CLASSES[code]
    oc = f"JSIC Rev. 13 {code}: {label}"
    est, emp, sales, emp_per, sales_per = r[5], r[6], r[7], r[8], r[9]
    base = {"officialClass": oc, "geo": JP, "source": SRC["t71"]}
    if trade:
        base["trade"] = trade
    if r[4] == "0_総数":
        rows.append(dict(base, item="business_count", value=est, unit="establishments", year="2021-06",
                         measure="private establishments (外国の会社及び法人でない団体を除く), 1 June 2021", **({"note": note} if note else {})))
        if isinstance(sales_per, (int, float)):
            rows.append(dict(base, item="sales_per_business", value=int(sales_per) * 10000, unit="JPY", year="2020",
                             measure="sales per establishment as published (1事業所当たり売上（収入）金額, in units of 10,000 JPY, here written in JPY), calendar year 2020",
                             note=" ".join(x for x in (note, f"Published figure: it divides sales by the establishments that reported sales, so it is higher than total sales ({sales:,} million JPY) over all {est:,} establishments.") if x)))
        if isinstance(emp_per, (int, float)):
            rows.append(dict(base, item="persons_per_business", value=emp_per, unit="persons engaged per establishment", year="2021-06",
                             measure="persons engaged per establishment as published (1事業所当たり従業者数), owners and family workers included",
                             **({"note": note} if note else {})))
    elif r[4] in BANDS and code != "AR":
        tot = None
    if r[4] in BANDS:
        rows.append(dict(base, item="size_class_share_n", value=est, unit="establishments", year="2021-06", measure=BANDS[r[4]]))

# turn the band counts into shares of the class total
tot = {(x["officialClass"]): x["value"] for x in rows if x["item"] == "business_count"}
out = []
for x in rows:
    if x["item"] == "size_class_share_n":
        n, t = x["value"], tot[x["officialClass"]]
        out.append(dict(x, item="size_class_share", value=round(100 * n / t, 1), unit="percent of establishments",
                        measure=f"share of private establishments with {x['measure']} persons engaged", note=f"{n:,} of {t:,} establishments."))
    else:
        out.append(x)
rows = out

# 2. Municipalities: table 9-1A (all establishments, private and public)
ws = openpyxl.load_workbook(f"{D}/000040067884.xlsx", read_only=True).active
idx = {}
for i, r in enumerate(ws.iter_rows(values_only=True)):
    if i == 5:
        for j, c in enumerate(r):
            if isinstance(c, str) and c.split("_")[0] in CLASSES:
                idx[c.split("_")[0]] = (j, c.split("_", 1)[1])
        continue
    if i < 9 or not isinstance(r[1], str):
        continue
    code = r[1][:5]
    if code not in CITIES and code != "00000":
        continue
    for k, (j, label) in idx.items():
        v = r[j]
        if not isinstance(v, (int, float)):
            continue
        trade, note = CLASSES[k]
        row = dict(item="business_count", officialClass=f"JSIC Rev. 13 {k}: {label}", geo=JP if code == "00000" else cg(code), value=v,
                   unit="establishments", year="2021-06",
                   measure="all establishments (全事業所数, table 9-1A), 1 June 2021: a wider count than the private establishments of table 7-1",
                   note=" ".join(x for x in (note, "Same table and definition for Japan and every city: compare cities on this count, not on table 7-1.") if x),
                   source=SRC["t91"])
        if trade:
            row["trade"] = trade
        rows.append(row)

# 3. Populations, 1 October 2020
ws = openpyxl.load_workbook(f"{D}/pop_000032142402.xlsx", read_only=True).active
for r in ws.iter_rows(values_only=True):
    if not r or not isinstance(r[5], str):
        continue
    code = r[5]
    if r[3] != "2000" and r[3] is not None:
        pass
    if code in CITIES or code == "00000":
        g = JP if code == "00000" else cg(code)
        if any(x["item"] == "population" and x["geo"]["code"] == g["code"] for x in rows):
            continue
        rows.append(dict(item="population", geo=g, value=int(r[7]), unit="persons", year="2020-10",
                         measure="usually resident population, Population Census of 1 October 2020",
                         note="Same municipal code as the business counts of table 9-1A (the census of businesses is eight months later).", source=SRC["pop"]))

write("estat-japan.json", {
    "id": "estat-japan",
    "source": {"publisher": "Statistics Bureau of Japan (Ministry of Internal Affairs and Communications) and Ministry of Economy, Trade and Industry",
               "title": "2021 Economic Census for Business Activity and 2020 Population Census, via e-Stat", "url": "https://www.e-stat.go.jp/stat-search/files?page=1&toukei=00200553&tstat=000001145590",
               "table": "Each row names its own table: Economic Census tables 7-1 and 9-1A, Population Census table 1-1",
               "licence": "e-Stat terms of use: Government of Japan Standard Terms of Use (version 2.0), compatible with CC BY 4.0, https://www.e-stat.go.jp/en/terms-of-use"},
    "retrieved": RETRIEVED, "rows": rows})

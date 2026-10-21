"""DESNZ table 3.4.2: prices of fuels purchased by non-domestic consumers in the UK, including the Climate Change Levy, excluding VAT."""
import os, sys, urllib.request
sys.path.insert(0, "/srv/atlas/wc/cwd/research-business/venv/lib/python3.13/site-packages")
import openpyxl
from spines_lib import DL, RETRIEVED, write

URL = "https://assets.publishing.service.gov.uk/media/6ab543da997a4b2950cceedc/table_341.xlsx"
path = os.path.join(DL, "uk", "table_341.xlsx")
if not os.path.exists(path):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    open(path, "wb").write(urllib.request.urlopen(urllib.request.Request(URL, headers={"User-Agent": "Mozilla/5.0"})).read())
wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
BANDS = {"Electricity: Very Small": ("energy_price_business", "electricity", "very small, 0 to 20 MWh a year"),
         "Electricity: Small": ("energy_price_business", "electricity", "small, 20 to 499 MWh a year"),
         "Electricity: Average": ("energy_price_business", "electricity", "all sizes, average"),
         "Gas: Very Small": ("gas_price_business", "gas", "very small, below 278 MWh a year"),
         "Gas: Small": ("gas_price_business", "gas", "small, 278 to 2,777 MWh a year"),
         "Gas: Average": ("gas_price_business", "gas", "all sizes, average")}
GB = {"level": "country", "code": "GB", "name": "United Kingdom"}
QN = {"1st": "Q1", "2nd": "Q2", "3rd": "Q3", "4th": "Q4"}
rows = []


def emit(head, r, year, period_words, prelim):
    for c, h in enumerate(head):
        if not h:
            continue
        key = str(h).replace(" (Pence per kWh)", "")
        if key not in BANDS or not isinstance(r[c], (int, float)):
            continue
        item, fuel, band = BANDS[key]
        rows.append({"item": item, "officialClass": f"{fuel}, non-domestic consumers, {band}", "geo": GB,
                     "value": round(r[c], 3), "unit": "GBP pence per kWh", "year": year,
                     "measure": f"average {fuel} price a kWh paid by non-domestic consumers, {band}, {period_words}",
                     "note": "Fully delivered price including the Climate Change Levy, excluding VAT (table 3.4.2)." + (" Provisional." if prelim else "")})


q = list(wb["3.4.2 (incl CCL)"].iter_rows(values_only=True))
head = q[7]
data = [r for r in q[8:] if isinstance(r[0], int) and r[1] in QN]
for r in data[-4:]:
    prelim = (r[0], r[1]) == (2026, "2nd")
    emit(head, r, f"{r[0]}-{QN[r[1]]}", f"{QN[r[1]]} {r[0]}", prelim)
a = list(wb["3.4.2 (Annual inc CCL)"].iter_rows(values_only=True))
ahead = a[6]
last = [r for r in a[7:] if isinstance(r[0], int)][-1]
emit(ahead, last, str(last[0]), f"calendar year {last[0]}", False)
write("business", "desnz-nondomestic-energy-prices.json", {
    "id": "desnz-nondomestic-energy-prices",
    "source": {"publisher": "Department for Energy Security and Net Zero", "title": "Gas and electricity prices in the non-domestic sector (Energy Prices, table 3.4.2)",
               "url": "https://www.gov.uk/government/statistical-data-sets/gas-and-electricity-prices-in-the-non-domestic-sector",
               "table": "table_341.xlsx (published 29/09/2026; Q2 2026 provisional), sheets '3.4.2 (incl CCL)' (the latest four quarters) and '3.4.2 (Annual inc CCL)' (the latest year); size bands from its Notes sheet",
               "licence": "Open Government Licence v3.0, https://www.nationalarchives.gov.uk/doc/open-government-licence/version/3/"},
    "retrieved": RETRIEVED, "rows": rows})

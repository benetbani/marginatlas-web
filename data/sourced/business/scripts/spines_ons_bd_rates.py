"""ONS Business demography 2024: enterprise birth and death rates per 100 active, derived from the release's own counts."""
import os, sys
sys.path.insert(0, "/srv/atlas/wc/cwd/research-business/venv/lib/python3.13/site-packages")
import openpyxl
from spines_lib import DL, RETRIEVED, write

wb = openpyxl.load_workbook(os.path.join(DL, "bd2024.xlsx"), read_only=True, data_only=True)


def by_sic(sheet):
    rows = list(wb[sheet].iter_rows(values_only=True))
    years = [str(y) for y in rows[3][1:7]]
    out = {}
    for r in rows[4:]:
        lab = str(r[0] or "").strip()
        key = "Total" if lab == "Total" else lab.split(" : ")[0]
        out[key] = dict(zip(years, r[1:7]))
    return out


def by_area(sheet):
    rows = list(wb[sheet].iter_rows(values_only=True))
    return {r[0]: r[2] for r in rows[4:] if r[0] in ("K02000001", "E12000007")}


B, D, A = by_sic("Table 1.2"), by_sic("Table 2.2"), by_sic("Table 3.2")
GROUPS = {"561": (["restaurants", "cafes"], "SIC 2007 group 561: Restaurants and mobile food service activities",
                  "Group 561 holds licensed restaurants (56.10/1), unlicensed restaurants and cafes (56.10/2), takeaways and mobile food together; the subclasses are not published."),
          "960": (["barbershops"], "SIC 2007 group 960: Other personal service activities",
                  "Group 960 holds hairdressing and beauty (96.02) with laundries, funeral and other personal services; 96.02 is not published."),
          "Total": (["all"], "SIC 2007 all industries", None)}
GB = {"level": "country", "code": "GB", "name": "United Kingdom"}
rows = []


def add(item, word, trade, cls, geo, year, num, den, note, table):
    notes = [n for n in (note,) if n]
    notes.append(f"Derived from the release's own counts: {num:,} enterprise {word} of {den:,} active enterprises (both control rounded to base 5; {table}).")
    if word == "deaths":
        notes.append("Deaths for the latest years are adjusted by the ONS for later re-activations (its Table 9).")
    rows.append({"item": item, "trade": trade, "officialClass": cls, "geo": geo,
                 "value": round(100 * num / den, 2), "unit": "per 100 active enterprises", "year": year,
                 "measure": f"enterprise {word} in {year} divided by active enterprises in {year}, percent",
                 "note": " ".join(notes)})


for g, (trades, cls, note) in GROUPS.items():
    for y in B[g]:
        for trade in trades:
            add("birth_rate", "births", trade, cls, GB, y, B[g][y], A[g][y], note, "Tables 1.2 and 3.2")
            add("death_rate", "deaths", trade, cls, GB, y, D[g][y], A[g][y], note, "Tables 2.2 and 3.2")
b, d, a = by_area("Table 1.1d"), by_area("Table 2.1d"), by_area("Table 3.1d")
LON = {"level": "region", "code": "E12000007", "name": "London"}
add("birth_rate", "births", "all", "SIC 2007 all industries", LON, "2024", b["E12000007"], a["E12000007"], "London region; no industry detail is published for a region.", "Tables 1.1d and 3.1d")
add("death_rate", "deaths", "all", "SIC 2007 all industries", LON, "2024", d["E12000007"], a["E12000007"], "London region; no industry detail is published for a region.", "Tables 2.1d and 3.1d")
write("business", "ons-bd-rates-2024.json", {
    "id": "ons-bd-rates-2024",
    "source": {"publisher": "Office for National Statistics", "title": "Business demography, UK: 2024 (reference tables)",
               "url": "https://www.ons.gov.uk/businessindustryandtrade/business/activitysizeandlocation/datasets/businessdemographyreferencetable",
               "table": "Tables 1.2, 2.2, 3.2 (births, deaths, active by SIC 2007 group, 2019 to 2024) and 1.1d, 2.1d, 3.1d (London region, 2024)",
               "licence": "Open Government Licence v3.0, https://www.nationalarchives.gov.uk/doc/open-government-licence/version/3/"},
    "retrieved": RETRIEVED, "rows": rows})

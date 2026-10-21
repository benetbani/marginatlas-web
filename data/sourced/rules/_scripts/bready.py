"""Extract B-READY 2025 rows for every covered economy into the sourced-data contract shape."""
import json, re, sys, unicodedata
import xml.etree.ElementTree as ET
import openpyxl, pycountry

ANS, SCORES, ISO, OUT = sys.argv[1:5]

def norm(s):
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode().upper()
    s = re.sub(r"\(THE\)", "", s)
    return re.sub(r"[^A-Z]", "", s)

# ISO 4217 list one: country name -> currency (first non-fund entry)
ccy = {}
for e in ET.parse(ISO).getroot().iter("CcyNtry"):
    n, c, f = e.findtext("CtryNm"), e.findtext("Ccy"), e.find("CcyNm")
    if not c or (f is not None and f.get("IsFund")):
        continue
    ccy.setdefault(norm(n), c)

# economy name -> ISO3 from the scores workbook
wb = openpyxl.load_workbook(SCORES, read_only=True)
iso3 = {}
for r in wb["00_B-READY_Pillar_Score"].iter_rows(min_row=2, values_only=True):
    if r[0] and r[1]:
        iso3[r[0].replace("\u2019", "'")] = r[1]

def a2_and_ccy(econ):
    c3 = iso3[econ]
    if c3 == "XKX":
        return "XK", "EUR", "Kosovo"
    alias = {"KOR": "KOREA (THE REPUBLIC OF)", "MDA": "MOLDOVA (THE REPUBLIC OF)"}
    if c3 in alias:
        return pycountry.countries.get(alpha_3=c3).alpha_2, ccy[norm(alias[c3])], econ
    p = pycountry.countries.get(alpha_3=c3)
    for n in (getattr(p, "official_name", None), p.name, getattr(p, "common_name", None)):
        if n and norm(n) in ccy:
            return p.alpha_2, ccy[norm(n)], p.name
    return p.alpha_2, None, p.name

WANT = {
    "BE_6_4_1": "entry", "BE_6_4_2": "entry_cost", "BE_2_1_1": "mincap",
    "bi_timedays_liq96": "liq_t", "bi_cost_liq96": "liq_c",
    "bi_timedays_liq97": "reo_t", "bi_cost_liq97": "reo_c", "tax7": "tax7",
}
ans = {}
qtext = {}
wb = openpyxl.load_workbook(ANS, read_only=True)
for ws in wb.worksheets[1:]:
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[1] in WANT and not (r[1] == "tax7" and ws.title.startswith("07")):
            ans.setdefault(r[0], {})[WANT[r[1]]] = r[3]
            qtext[WANT[r[1]]] = str(r[2]).strip()

CITY = "Measured in the economy's largest business city (B-READY Methodology Handbook, 2nd edition); the workbook does not name the city."
# Bulgaria adopted the euro on 1 January 2026; B-READY 2025 answers are as of 1 September 2024, in the currency of that date.
CCY_AT_2024 = {"BG": "BGN"}

def num(v):
    try:
        x = float(v)
    except (TypeError, ValueError):
        return None
    return None if x == -777 else (int(x) if x == int(x) else round(x, 4))

rows, missing = [], []
for econ in sorted(ans):
    a = ans[econ]
    code, cur, pname = a2_and_ccy(econ)
    cur = CCY_AT_2024.get(code, cur)
    if cur is None:
        missing.append(econ)
    geo = {"level": "country", "code": code, "name": econ}
    def row(item, key, unit, measure, extra=""):
        raw = a.get(key)
        v = num(raw)
        r = {"item": item, "geo": geo, "value": v, "unit": unit, "year": "2025", "measure": measure,
             "note": (CITY + (" " + extra if extra else "")).strip()}
        if v is None:
            r["note"] = (f"The workbook answers '{raw}' (-777: no practice in the area measured or no data). " + r["note"]) if raw is not None else ("The workbook holds no answer for this economy. " + r["note"])
        rows.append(r)
    row("company_entry_days", "entry", "calendar days", qtext["entry"])
    extra_cost = {"GB": "The 50 is the Companies House online incorporation fee before it rose to GBP 100 on 1 February 2026.",
                  "BG": "Read as lev: the workbook (23 December 2025) predates Bulgaria's adoption of the euro on 1 January 2026, and the answers were collected before it."}.get(code, "")
    row("company_entry_cost", "entry_cost", f"{cur} (local currency units)" if cur else "local currency units", qtext["entry_cost"],
        ("The handbook collects this cost in local currency units. " + extra_cost).strip() + ("" if cur else " The ISO 4217 list names no single currency for this economy, and the workbook does not name it."))
    m = a.get("mincap")
    mv = 1 if m == "Yes" else 0 if m == "No" else None
    r = {"item": "minimum_capital_required", "geo": geo, "value": mv, "unit": "1 = yes, 0 = no", "year": "2025",
         "measure": qtext["mincap"], "note": CITY}
    if mv is None:
        r["note"] = f"The workbook answers '{m}'. " + CITY
    rows.append(r)
    row("liquidation_months", "liq_t", "months", qtext["liq_t"],
        "The handbook asks for calendar months, from filing until creditors are repaid all or some of what they are owed.")
    row("liquidation_cost_pct", "liq_c", "% of the value of the company's assets", qtext["liq_c"],
        "Court, administrator, lawyer and other fees, as a percentage of the company's asset value (handbook).")
    row("reorganisation_months", "reo_t", "months", qtext["reo_t"],
        "From filing until creditors approve the reorganisation plan (handbook). The handbook states calendar months for the liquidation question; the reorganisation question names no unit and its answers are read in the same unit.")
    row("reorganisation_cost_pct", "reo_c", "% of the value of the company's assets", qtext["reo_c"],
        "Court, administrator, lawyer and other fees, as a percentage of the company's asset value (handbook).")
    r_ = None
    v = num(a.get("tax7"))
    rows.append({"item": "labour_tax_share_of_labour_cost", "geo": geo, "value": None if v is None else round(v, 1),
                 "unit": "% of annual cost of labour", "year": "2025",
                 "measure": "Reported share of social security and employment-based taxes in annual cost of labor (firms' own answers, World Bank Enterprise Surveys)",
                 "note": ("What firms report, not a statutory rate. Total annual social security payments and employment-based taxes (employee taxes withheld excluded), divided by total annual cost of labour, last fiscal year (handbook)."
                          if v is not None else f"The workbook answers '{a.get('tax7')}' (-777: indicator not available from the Enterprise Surveys).")})

out = {
    "id": "b-ready-2025",
    "source": {"publisher": "World Bank", "title": "Business Ready 2025, economy-level answers (B-READY_ALL_DATA_2025.zip, workbook 02_B-READY-2025-EconomyAnswer.xlsx, file version 23 December 2025)",
               "url": "https://www.worldbank.org/content/dam/sites/b-ready/documents/excel/B-READY_ALL_DATA_2025.zip",
               "table": "01_Business_Entry Q25, Q119, Q120; 04_Labor tax7; 10_Business_Insolvency Q64 to Q67"},
    "retrieved": "2026-10-17",
    "rows": rows,
}
import re
for r in rows:
    for k in ("note", "measure"):
        if isinstance(r.get(k), str):
            r[k] = re.sub(r";\s+(\S)", lambda m: ". " + m.group(1).upper(), r[k])
json.dump(out, open(OUT, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
print("economies", len(ans), "rows", len(rows), "no currency:", missing)
